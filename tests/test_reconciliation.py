"""Behavioral regression tests using only invented shipments and prices.

Expected amounts below are hand-calculated from the invented tariff, rather
than obtained from implementation helpers. All workbooks live in temporary
directories, so running this suite never requires or changes business data.
"""

from __future__ import annotations

import copy
import datetime as dt
import json
from pathlib import Path
import sys
import tempfile
import unittest
from decimal import Decimal

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))


COURIER_HEADERS = ["物流单号", "重量", "揽收目的地网点", "日期", "差额", "充单", "运费"]
SYSTEM_HEADERS = ["发货时间", "物流单号", "物流公司", "商品编码", "商品件数", "重量", "运费", "省份", "城市", "货区"]
PRICE_HEADERS = ["省份", "0-0.3kg", "0.3-0.5kg", "0.5-1kg", "1-2kg", "2-3kg", "3-5kg", "5-8kg", "8-10kg", "10-15kg", "15-20kg", "20-25kg", "30kg"]
INVENTED_PRICES = [2, 3, 5, 8, 8, 10, 12, 14, 16, 18, 20, 22]
DATE = dt.datetime(2040, 9, 10)


def courier_row(number, weight, charge, *, address="浙江省杭州市", date=DATE):
    return {"物流单号": number, "重量": weight, "揽收目的地网点": address,
            "日期": date, "差额": charge, "充单": 0, "运费": charge}


def system_row(number, weight, *, sku="虚构商品*2", address="浙江省杭州市", date=DATE, freight=999):
    return {"发货时间": date, "物流单号": number, "物流公司": "圆通速递",
            "商品编码": sku, "商品件数": 1, "重量": weight, "运费": freight,
            "省份": address, "城市": "", "货区": ""}


def write_table(path, sheet, headers, records):
    workbook = Workbook()
    ws = workbook.active
    ws.title = sheet
    ws.append(headers)
    for record in records:
        ws.append([record.get(header) for header in headers])
    workbook.save(path)


def write_prices(path, *, second_version=False):
    workbook = Workbook()
    ws = workbook.active
    ws.title = "虚构报价"
    versions = ["9月1日-15日价格", "9月16日-30日价格"] if second_version else ["9月1日-30日价格"]
    for title in versions:
        ws.append([title])
        ws.append(PRICE_HEADERS)
        for region in ["浙江", "北京", "广东", "深圳", "甘肃"]:
            ws.append([region, *INVENTED_PRICES])
        ws.append([])
    workbook.save(path)


class ReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="reconciliation-test-")
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        self.courier = self.directory / "courier.xlsx"
        self.system = self.directory / "system.xlsx"
        self.price = self.directory / "price.xlsx"
        self.profile = json.loads((ROOT / "profiles" / "yuantong-jushuitan-v1.json").read_text(encoding="utf-8"))

    def inputs(self, courier, system, *, second_version=False):
        write_table(self.courier, "快递明细", COURIER_HEADERS, courier)
        write_table(self.system, "系统明细", SYSTEM_HEADERS, system)
        write_prices(self.price, second_version=second_version)

    def run_reconciliation(self, profile=None):
        from reconciliation_core import reconcile
        return reconcile(self.courier, self.system, self.price, profile or self.profile)

    @staticmethod
    def rows(result, section="all_records"):
        return {row["物流单号"]: row for row in result[section] if row.get("物流单号")}

    @staticmethod
    def difference(row):
        return row.get("总差额", row.get("差额"))

    def assert_money(self, row, *, billed, courier, system, total, weight, pricing):
        expectations = {"快递账单收费": billed, "按快递重量计算": courier,
                        "按系统重量计算": system, "重量影响金额": weight,
                        "计价差异金额": pricing}
        for field, expected in expectations.items():
            actual = row[field]
            self.assertIsInstance(actual, Decimal, field)
            self.assertEqual(actual, Decimal(expected), field)
        self.assertEqual(self.difference(row), Decimal(total))
        self.assertEqual(self.difference(row), row["重量影响金额"] + row["计价差异金额"])

    def scenario(self):
        # id, courier kg, system kg, billed yuan; prices are above, not engine-derived.
        cases = [
            ("FAKE-SAME", .29, .2, 2),
            ("FAKE-WEIGHT", .4, .2, 3),
            ("FAKE-PRICE", .2, .2, 3),
            ("FAKE-MIX-POS", .4, .2, 2.5),
            ("FAKE-MIX-ZERO", .4, .2, 2),
            ("FAKE-NEG", .2, .4, 2),
            ("FAKE-CROSS-ZERO", 2.1, 1.9, 8),
            ("FAKE-FINE", 6, 4, 12),
            ("FAKE-FINE-SAME", 4.5, 4, 10),
        ]
        self.inputs([courier_row(number, cw, fee) for number, cw, sw, fee in cases],
                    [system_row(number, sw) for number, cw, sw, fee in cases])

    def test_hand_calculated_causes_totals_and_positive_only_evidence(self):
        self.scenario()
        result = self.run_reconciliation()
        rows = self.rows(result)
        expected = {
            "FAKE-WEIGHT": ("重量影响运费", "3", "3", "2", "1", "1", "0"),
            "FAKE-PRICE": ("计价不符", "3", "2", "2", "1", "0", "1"),
            "FAKE-MIX-POS": ("重量＋计价差异", "2.5", "3", "2", ".5", "1", "-.5"),
            "FAKE-MIX-ZERO": ("重量＋计价差异", "2", "3", "2", "0", "1", "-1"),
            "FAKE-NEG": ("重量影响运费", "2", "2", "3", "-1", "-1", "0"),
            "FAKE-CROSS-ZERO": ("跨档但运费不变", "8", "8", "8", "0", "0", "0"),
            "FAKE-FINE": ("重量影响运费", "12", "12", "10", "2", "2", "0"),
        }
        for number, values in expected.items():
            with self.subTest(number=number):
                reason, billed, courier, system, total, weight, pricing = values
                self.assertEqual(rows[number]["主原因"], reason)
                self.assert_money(rows[number], billed=billed, courier=courier, system=system,
                                  total=total, weight=weight, pricing=pricing)
        self.assertEqual(set(self.rows(result, "details")), set(expected))
        self.assertEqual(set(self.rows(result, "evidence")),
                         {"FAKE-WEIGHT", "FAKE-PRICE", "FAKE-MIX-POS", "FAKE-FINE"})
        self.assertEqual(sum((self.difference(r) for r in result["details"]), Decimal(0)), Decimal("3.50"))
        self.assertEqual(sum((self.difference(r) for r in result["evidence"]), Decimal(0)), Decimal("4.50"))
        self.assertEqual(len(result["products"]), 1)
        product = result["products"][0]
        self.assertEqual(product["商品编码"], "虚构商品*2")
        self.assertEqual(product["异常总单量"], 7)
        self.assertEqual(product["总差额"], Decimal("3.50"))
        self.assertEqual(product["账单偏高金额"], Decimal("4.50"))
        self.assertEqual(product["账单偏低金额"], Decimal("1.00"))

    def test_identical_inputs_have_identical_semantic_result(self):
        self.scenario()
        before = {p: p.read_bytes() for p in [self.courier, self.system, self.price]}
        first = self.run_reconciliation()
        second = self.run_reconciliation()
        self.assertEqual(first, second)
        self.assertEqual(before, {p: p.read_bytes() for p in before})

    def test_system_freight_is_ignored_even_formula_missing_or_extreme(self):
        self.scenario()
        original = self.run_reconciliation()

        def business(result):
            fields = ["物流单号", "商品编码", "主原因", "快递重量kg", "系统重量kg",
                      "快递账单收费", "按快递重量计算", "按系统重量计算",
                      "重量影响金额", "计价差异金额"]
            return {section: [(tuple(row.get(k) for k in fields), self.difference(row))
                              for row in result[section]]
                    for section in ["all_records", "details", "evidence"]}

        for value in [None, -999999999, 987654321, "=1/0", "=SUM(A1:Z99)", "任意非金额"]:
            with self.subTest(system_freight=value):
                workbook = load_workbook(self.system)
                ws = workbook.active
                for row in range(2, ws.max_row + 1):
                    ws.cell(row, 7).value = value
                workbook.save(self.system)
                changed = self.run_reconciliation()
                self.assertEqual(business(original), business(changed))
                self.assertEqual(original["products"], changed["products"])
        workbook = load_workbook(self.system)
        workbook.active.delete_cols(7)
        workbook.save(self.system)
        changed = self.run_reconciliation()
        self.assertEqual(business(original), business(changed))

    def test_region_and_date_uncertainty_are_not_quantified_as_overcharge(self):
        self.inputs([
            courier_row("FAKE-REGION", .4, 90, address="浙江省杭州市"),
            courier_row("FAKE-DATE", .4, 90, date=dt.datetime(2040, 9, 15)),
            courier_row("FAKE-NO-RATE", .4, 90, address="海南省海口市"),
        ], [
            system_row("FAKE-REGION", .2, address="北京市"),
            system_row("FAKE-DATE", .2, date=dt.datetime(2040, 9, 16)),
            system_row("FAKE-NO-RATE", .2, address="海南省海口市"),
        ], second_version=True)
        result = self.run_reconciliation()
        rows = self.rows(result)
        self.assertEqual(rows["FAKE-REGION"]["主原因"], "地区待确认")
        self.assertEqual(rows["FAKE-DATE"]["主原因"], "日期待确认")
        self.assertIsNone(self.difference(rows["FAKE-REGION"]))
        self.assertIsNone(self.difference(rows["FAKE-DATE"]))
        self.assertEqual(result["evidence"], [])
        self.assertEqual(rows["FAKE-NO-RATE"]["主原因"], "报价规则待确认")
        self.assertEqual(result["summary"]["未知差额单量"], 3)
        for field in ["总差额", "账单偏高金额", "账单偏低金额"]:
            self.assertIsNone(result["summary"][field], field)
        issue_count = sum(issue.get("严重程度") != "说明" for issue in result["rule_issues"])
        self.assertEqual(result["summary"]["价格表问题数"], issue_count)

    def test_duplicate_is_not_first_row_wins_and_unmatched_courier_still_prices(self):
        self.inputs([courier_row("FAKE-DUP", .4, 3), courier_row("FAKE-ONLY-C", .4, 4)],
                    [system_row("FAKE-DUP", .2), system_row("FAKE-DUP", .4),
                     system_row("FAKE-ONLY-S", .2)])
        result = self.run_reconciliation()
        rows = self.rows(result)
        self.assertEqual(rows["FAKE-DUP"]["主原因"], "关联待核实")
        self.assertIsNone(rows["FAKE-DUP"]["按系统重量计算"])
        self.assertEqual(rows["FAKE-DUP"]["按快递重量计算"], Decimal("3.00"))
        self.assertEqual(rows["FAKE-ONLY-C"]["主原因"], "仅快递有单")
        self.assertEqual(rows["FAKE-ONLY-C"]["按快递重量计算"], Decimal("3.00"))
        self.assertEqual(rows["FAKE-ONLY-C"]["计价差异金额"], Decimal("1.00"))
        self.assertIsNone(self.difference(rows["FAKE-ONLY-C"]))
        self.assertEqual(rows["FAKE-ONLY-S"]["主原因"], "仅系统有单")
        self.assertEqual(result["evidence"], [])

    def test_text_ids_preserve_zeros_and_numeric_precision_loss_is_not_guessed(self):
        self.inputs([courier_row("  000FAKE001  ", .2, 2),
                     courier_row("00001234567890123456", .2, 2),
                     courier_row(1234567890123456, .2, 90),
                     courier_row(None, .2, 90)],
                    [system_row("000FAKE001", .2),
                     system_row("00001234567890123456", .2)])
        result = self.run_reconciliation()
        rows = self.rows(result)
        self.assertIn("000FAKE001", rows)
        self.assertIn("00001234567890123456", rows)
        self.assertEqual(rows["00001234567890123456"]["按系统重量计算"], Decimal("2.00"))
        self.assertNotIn("1234567890123456", self.rows(result, "evidence"))
        serialized = json.dumps(result, ensure_ascii=False, default=str)
        self.assertIn("精度", serialized)
        self.assertTrue("空单号" in serialized or "单号为空" in serialized or "缺少单号" in serialized)

    def test_weight_boundaries_are_not_rounded_or_silently_extended(self):
        cases = [("FAKE-ZERO", 0, 2), ("FAKE-LOW-END", .3, 2),
                 ("FAKE-NEXT-START", .301, 3), ("FAKE-NEXT-END", .5, 3),
                 ("FAKE-GAP", .3005, 90), ("FAKE-HIGH-EDGE", 5, 90),
                 ("FAKE-LAST-UNKNOWN", 26, 90), ("FAKE-OUTSIDE", 31, 90)]
        self.inputs([courier_row(n, w, a) for n, w, a in cases],
                    [system_row(n, w) for n, w, a in cases])
        result = self.run_reconciliation()
        rows = self.rows(result)
        for number, amount in [("FAKE-ZERO", "2"), ("FAKE-LOW-END", "2"),
                               ("FAKE-NEXT-START", "3"), ("FAKE-NEXT-END", "3")]:
            self.assertEqual(rows[number]["按系统重量计算"], Decimal(amount))
        for number in ["FAKE-GAP", "FAKE-HIGH-EDGE", "FAKE-LAST-UNKNOWN", "FAKE-OUTSIDE"]:
            self.assertIsNone(rows[number]["按系统重量计算"], number)
            self.assertIsNone(self.difference(rows[number]), number)
            self.assertNotIn(number, self.rows(result, "evidence"))

    def test_column_reordering_and_documented_aliases_preserve_amounts(self):
        self.inputs([courier_row("FAKE-ALIAS", .4, 3)], [system_row("FAKE-ALIAS", .2)])
        original = self.run_reconciliation()
        for path, aliases in [(self.courier, {"物流单号": "运单号", "重量": "计费重量", "运费": "账单运费"}),
                              (self.system, {"物流单号": "快递单号", "重量": "实重", "商品编码": "SKU"})]:
            workbook = load_workbook(path)
            ws = workbook.active
            values = list(ws.values)
            headers = list(reversed(values[0]))
            reordered = [[aliases.get(h, h) for h in headers], *[list(reversed(row)) for row in values[1:]]]
            ws.delete_rows(1, ws.max_row)
            for row in reordered:
                ws.append(row)
            workbook.save(path)
        row = self.rows(self.run_reconciliation())["FAKE-ALIAS"]
        self.assert_money(row, billed="3", courier="3", system="2", total="1", weight="1", pricing="0")
        self.assertEqual(row["主原因"], self.rows(original)["FAKE-ALIAS"]["主原因"])

    def test_multiple_candidate_sheets_require_explicit_selection(self):
        from reconciliation_core import ReconciliationError
        self.inputs([courier_row("FAKE-SHEET", .2, 2)], [system_row("FAKE-SHEET", .2)])
        workbook = load_workbook(self.courier)
        workbook.copy_worksheet(workbook.active).title = "另一个有效账单"
        workbook.save(self.courier)
        with self.assertRaises(ReconciliationError):
            self.run_reconciliation()
        profile = copy.deepcopy(self.profile)
        profile["sheets"]["courier"] = "快递明细"
        result = self.run_reconciliation(profile)
        self.assertEqual(set(self.rows(result)), {"FAKE-SHEET"})

    def test_complete_address_city_exception_and_province_prefix(self):
        self.inputs([courier_row("FAKE-CITY", .2, 3, address="广东省深圳市光明区"),
                     courier_row("FAKE-PREFIX", .2, 2, address="甘肃省兰州市黄河北路")],
                    [system_row("FAKE-CITY", .2, address="广东省 深圳市 光明区"),
                     system_row("FAKE-PREFIX", .2, address="甘肃省兰州市")])
        workbook = load_workbook(self.price)
        ws = workbook.active
        for row in ws:
            if row[0].value == "深圳":
                ws.cell(row[0].row, 2).value = 3
        workbook.save(self.price)
        rows = self.rows(self.run_reconciliation())
        self.assertEqual(rows["FAKE-CITY"]["按系统重量计算"], Decimal("3.00"))
        self.assertEqual(rows["FAKE-PREFIX"]["按系统重量计算"], Decimal("2.00"))
        self.assertEqual(self.difference(rows["FAKE-CITY"]), Decimal("0.00"))

    def test_known_formula_recalculated_but_unknown_formula_not_executed(self):
        a = courier_row("FAKE-FORMULA", .4, 3)
        a["运费"] = "=E2+F2"
        b = courier_row("FAKE-UNKNOWN-FORMULA", .4, 3)
        b["运费"] = "=SUM(E3:F3)"
        self.inputs([a, b], [system_row("FAKE-FORMULA", .2), system_row("FAKE-UNKNOWN-FORMULA", .2)])
        result = self.run_reconciliation()
        rows = self.rows(result)
        self.assertEqual(rows["FAKE-FORMULA"]["快递账单收费"], Decimal("3.00"))
        self.assertEqual(self.difference(rows["FAKE-FORMULA"]), Decimal("1.00"))
        self.assertIsNone(rows["FAKE-UNKNOWN-FORMULA"]["快递账单收费"])
        self.assertIsNone(self.difference(rows["FAKE-UNKNOWN-FORMULA"]))
        self.assertNotIn("FAKE-UNKNOWN-FORMULA", self.rows(result, "evidence"))

    def test_half_up_cent_rounding_is_fixed(self):
        self.inputs([courier_row("FAKE-ROUND", .2, 2.015)], [system_row("FAKE-ROUND", .2)])
        row = self.rows(self.run_reconciliation())["FAKE-ROUND"]
        self.assertEqual(row["快递账单收费"], Decimal("2.02"))
        self.assertEqual(self.difference(row), Decimal("0.02"))

    def test_price_rule_conflict_does_not_select_first_or_cheapest_rule(self):
        self.inputs([courier_row("FAKE-CONFLICT", .2, 90)], [system_row("FAKE-CONFLICT", .2)])
        workbook = load_workbook(self.price)
        workbook.active.append(["浙江", 1, *INVENTED_PRICES[1:]])
        workbook.save(self.price)
        result = self.run_reconciliation()
        row = self.rows(result)["FAKE-CONFLICT"]
        self.assertIsNone(row["按系统重量计算"])
        self.assertIsNone(self.difference(row))
        self.assertTrue(result["rule_issues"])
        self.assertEqual(result["evidence"], [])

    def test_missing_quote_only_blocks_affected_weight(self):
        self.inputs([courier_row("FAKE-MISSING-RATE", .2, 90),
                     courier_row("FAKE-GOOD-RATE", .4, 3)],
                    [system_row("FAKE-MISSING-RATE", .2), system_row("FAKE-GOOD-RATE", .4)])
        workbook = load_workbook(self.price)
        ws = workbook.active
        for cells in ws:
            if cells[0].value == "浙江":
                ws.cell(cells[0].row, 2).value = None
        workbook.save(self.price)
        result = self.run_reconciliation()
        rows = self.rows(result)
        self.assertIsNone(rows["FAKE-MISSING-RATE"]["按系统重量计算"])
        self.assertEqual(rows["FAKE-GOOD-RATE"]["按系统重量计算"], Decimal("3.00"))
        self.assertEqual(self.difference(rows["FAKE-GOOD-RATE"]), Decimal("0.00"))
        self.assertTrue(result["rule_issues"])

    def test_missing_system_weight_does_not_suppress_courier_calculation(self):
        self.inputs([courier_row("FAKE-MISSING-WEIGHT", .4, 4)],
                    [system_row("FAKE-MISSING-WEIGHT", None)])
        row = self.rows(self.run_reconciliation())["FAKE-MISSING-WEIGHT"]
        self.assertEqual(row["按快递重量计算"], Decimal("3.00"))
        self.assertEqual(row["计价差异金额"], Decimal("1.00"))
        self.assertIsNone(row["按系统重量计算"])
        self.assertIsNone(self.difference(row))

    def test_system_freight_only_row_does_not_create_a_record(self):
        self.inputs([courier_row("FAKE-EMPTY-ROW", .2, 2)], [system_row("FAKE-EMPTY-ROW", .2)])
        # Keep the physical blank row in every variant so source-row metadata
        # cannot change merely because the worksheet used range grew.
        workbook = load_workbook(self.system)
        workbook.active.cell(3, 1).number_format = "@"
        workbook.save(self.system)
        baseline = self.run_reconciliation()
        for value in [1234567, "=1/0", "仅运费有值", None]:
            with self.subTest(value=value):
                workbook = load_workbook(self.system)
                workbook.active.cell(3, 7).value = value
                workbook.save(self.system)
                result = self.run_reconciliation()
                for section in ["all_records", "details", "products", "evidence", "summary"]:
                    self.assertEqual(result[section], baseline[section], section)

    def test_system_freight_cannot_be_remapped_as_weight(self):
        from reconciliation_core import ReconciliationError
        self.inputs([courier_row("FAKE-BAD-MAPPING", .2, 2)], [system_row("FAKE-BAD-MAPPING", .2)])
        profile = copy.deepcopy(self.profile)
        profile["aliases"]["system"]["weight"] = ["运费"]
        with self.assertRaises(ReconciliationError):
            self.run_reconciliation(profile)

    def test_zero_price_is_a_flagged_suspicion_but_remains_zero(self):
        self.inputs([courier_row("FAKE-ZERO-PRICE", .2, 0)], [system_row("FAKE-ZERO-PRICE", .2)])
        workbook = load_workbook(self.price)
        for cells in workbook.active:
            if cells[0].value == "浙江":
                workbook.active.cell(cells[0].row, 2).value = 0
        workbook.save(self.price)
        result = self.run_reconciliation()
        row = self.rows(result)["FAKE-ZERO-PRICE"]
        self.assertEqual(row["按快递重量计算"], Decimal("0.00"))
        self.assertEqual(row["按系统重量计算"], Decimal("0.00"))
        self.assertEqual(self.difference(row), Decimal("0.00"))
        self.assertTrue(result["rule_issues"])
        self.assertEqual(result["evidence"], [])

    def test_equal_price_boundary_retains_amount_but_waits_for_rule_confirmation(self):
        self.inputs([courier_row("FAKE-EDGE-C", 5, 11), courier_row("FAKE-EDGE-S", 4, 11)],
                    [system_row("FAKE-EDGE-C", 4), system_row("FAKE-EDGE-S", 5)])
        workbook = load_workbook(self.price)
        ws = workbook.active
        for cells in ws:
            if cells[0].value == "浙江":
                # Invented adjacent tariffs: both 3–5kg and 5–8kg cost 10 yuan.
                ws.cell(cells[0].row, 8).value = 10
        workbook.save(self.price)
        result = self.run_reconciliation()
        for number in ["FAKE-EDGE-C", "FAKE-EDGE-S"]:
            row = self.rows(result)[number]
            self.assert_money(row, billed="11", courier="10", system="10",
                              total="1", weight="0", pricing="1")
            self.assertEqual(row["主原因"], "报价规则待确认")
        self.assertEqual(len(result["details"]), 2)
        self.assertEqual(result["evidence"], [])

    def test_explicit_carrier_policy_accepts_same_brand_aliases_but_rejects_other_brands(self):
        from reconciliation_core import ReconciliationError
        courier = [courier_row("FAKE-BRAND-A", .2, 2), courier_row("FAKE-BRAND-B", .2, 2)]
        system = [system_row("FAKE-BRAND-A", .2), system_row("FAKE-BRAND-B", .2)]
        system[0]["物流公司"] = "虚构甲网点-圆通速递"
        system[1]["物流公司"] = "虚构乙网点-圆通快递_77"
        self.inputs(courier, system)
        profile = copy.deepcopy(self.profile)
        profile["carrier_policy"] = "yuantong_brand_aliases"
        result = self.run_reconciliation(profile)
        self.assertEqual(set(self.rows(result)), {"FAKE-BRAND-A", "FAKE-BRAND-B"})
        for row in result["all_records"]:
            self.assertEqual(row["按系统重量计算"], Decimal("2.00"))
        system[1]["物流公司"] = "虚构乙网点-其他品牌快递"
        write_table(self.system, "系统明细", SYSTEM_HEADERS, system)
        with self.assertRaises(ReconciliationError):
            self.run_reconciliation(profile)

    def test_confirmed_upper_endpoints_use_the_lower_band_including_exact_three_kg(self):
        # Give every neighboring low-weight band a different invented amount,
        # including 2–3kg=9 and 3–5kg=10, so an off-by-one band cannot pass.
        cases = [("FAKE-END-03", .3, .3, 2), ("FAKE-END-05", .5, .5, 3),
                 ("FAKE-END-1", 1, 1, 5), ("FAKE-END-2", 2, 2, 8),
                 ("FAKE-END-3", 3, 3, 9), ("FAKE-END-3-SAME", 3, 2.99, 9)]
        self.inputs([courier_row(number, cw, fee) for number, cw, sw, fee in cases],
                    [system_row(number, sw) for number, cw, sw, fee in cases])
        workbook = load_workbook(self.price)
        for cells in workbook.active:
            if cells[0].value == "浙江":
                workbook.active.cell(cells[0].row, 6).value = 9
        workbook.save(self.price)
        result = self.run_reconciliation()
        rows = self.rows(result)
        for number, cw, sw, fee in cases:
            with self.subTest(number=number):
                expected = str(fee)
                self.assert_money(rows[number], billed=expected, courier=expected,
                                  system=expected, total="0", weight="0", pricing="0")
                self.assertEqual(rows[number]["主原因"], "")
        self.assertEqual(result["details"], [])
        self.assertEqual(result["summary"]["同档小差忽略单量"], 1)


if __name__ == "__main__":
    unittest.main()
