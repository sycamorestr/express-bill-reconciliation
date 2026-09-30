"""Deterministic, read-only freight reconciliation; system freight is never used.

Public API: reconcile(courier: Path, system: Path, price: Path, profile: dict).
Decimal amounts are retained; dates and source locations are plain strings.
"""
from __future__ import annotations

import collections
import datetime as dt
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
import json
from pathlib import Path
import re

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.datetime import from_excel

ENGINE_VERSION = "1.1.0"
PROFILE_PATH = Path(__file__).resolve().parent.parent / "profiles" / "yuantong-jushuitan-v1.json"
D = Decimal
CENT = D("0.01")
ZERO = D("0.00")
PROVINCES = ("黑龙江", "内蒙古", "新疆", "西藏", "广西", "宁夏", "北京", "天津", "上海", "重庆", "河北", "河南", "山东", "山西", "安徽", "江苏", "浙江", "湖北", "湖南", "福建", "江西", "广东", "海南", "四川", "贵州", "云南", "陕西", "甘肃", "青海", "辽宁", "吉林", "台湾", "香港", "澳门")
BANDS = ((D("0"), D(".3"), "0.3kg重量"), (D(".301"), D(".5"), "0.5kg重量"), (D(".501"), D("1"), "1kg重量"), (D("1.001"), D("2"), "2kg重量"), (D("2.001"), D("3"), "3kg重量"))
LOW_INTERVALS = {(D("0"), D(".3")): BANDS[0], (D(".3"), D(".5")): BANDS[1], (D(".5"), D("1")): BANDS[2], (D("1"), D("2")): BANDS[3], (D("2"), D("3")): BANDS[4]}


def default_profile():
    return json.loads(PROFILE_PATH.read_text(encoding="utf-8"))


class ReconciliationError(ValueError):
    """Unsafe or ambiguous input configuration requiring an explicit correction."""


def reconcile(courier: Path, system: Path, price: Path, profile: dict) -> dict:
    with localcontext() as context:
        context.prec = 50
        context.rounding = ROUND_HALF_UP
        return _reconcile(Path(courier), Path(system), Path(price), profile)


def _text(value):
    return "" if value is None else str(value).strip()


def _decimal(value):
    if value is None or isinstance(value, (bool, dt.date)):
        return None
    try:
        number = D(str(value).strip())
        return number if number.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def _money(value):
    number = _decimal(value)
    try:
        return number.quantize(CENT, rounding=ROUND_HALF_UP) if number is not None else None
    except InvalidOperation:
        return None


def _date(value, epoch):
    try:
        if isinstance(value, bool) or value is None:
            return None
        if isinstance(value, (int, float, D)):
            value = from_excel(float(value), epoch)
        if isinstance(value, dt.datetime):
            return value.date()
        if isinstance(value, dt.date):
            return value
        value = _text(value).replace("/", "-")
        # Whitelist extended calendar dates across Python versions. Python 3.11
        # broadened fromisoformat; compact/week dates must not change our output.
        simple = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T](\d{2}):(\d{2})(?::(\d{2})(?:\.(\d{1,6}))?)?(Z|[+-]\d{2}:\d{2})?)?", value)
        if not simple:
            return None
        year, month, day, hour, minute, second, fraction, offset = simple.groups()
        parsed_date = dt.date(int(year), int(month), int(day))
        if hour is None:
            return parsed_date
        normalized = f"{parsed_date.isoformat()}T{hour}:{minute}:{second or '00'}"
        if fraction:
            normalized += "." + fraction
        if offset:
            normalized += "+00:00" if offset == "Z" else offset
        return dt.datetime.fromisoformat(normalized).date()
    except (ValueError, TypeError, OverflowError):
        return None


def _band(weight):
    if weight is None or weight < 0:
        return "重量数据异常"
    if weight > 3:
        return "3kg以上重量"
    return next((name for low, high, name in BANDS if low <= weight <= high), "重量边界待确认")


def _identifier(value):
    if value is None or not _text(value):
        return None, "空单号"
    if isinstance(value, bool):
        return None, "单号数据类型异常"
    if isinstance(value, (int, float, D)):
        number = _decimal(value)
        if number is None or number < 0 or number != number.to_integral_value():
            return None, "单号精度待核实"
        text = format(number, "f").split(".")[0]
        if len(text) > 15:
            return None, "单号精度待核实"
        return text, ""
    text = _text(value)
    if re.fullmatch(r"[+-]?\d+(?:\.\d+)?[eE][+-]?\d+", text) or re.fullmatch(r"\d+\.\d+", text):
        return None, "单号精度待核实"
    if text.startswith("="):
        return None, "公式单号待核实"
    return text, ""


def _header(value):
    return re.sub(r"\s+", "", _text(value)).lower()


def _select_sheet(book, side, profile):
    explicit = profile.get("sheets", {}).get(side)
    explicit_row = profile.get("header_rows", {}).get(side)
    names = [explicit] if explicit else book.sheetnames
    if explicit and explicit not in book.sheetnames:
        raise ReconciliationError(f"{side} 指定工作表不存在：{explicit}")
    aliases = profile["aliases"][side]
    candidates = []
    for name in names:
        sheet = book[name]
        start = int(explicit_row or 1)
        stop = int(explicit_row or profile.get("header_search_rows", 30))
        for values in sheet.iter_rows(min_row=start, max_row=stop):
            mapping = {}
            ambiguities = []
            for field, options in aliases.items():
                positions = [cell.column for cell in values if _header(cell.value) in {_header(x) for x in options}]
                if len(positions) > 1:
                    ambiguities.append(field)
                elif positions:
                    mapping[field] = positions[0]
            if "id" in mapping or "id" in ambiguities:
                candidates.append((name, values[0].row, mapping, ambiguities))
    if len(candidates) != 1:
        positions = ", ".join(f"{x[0]}!{x[1]}" for x in candidates) or "无"
        raise ReconciliationError(f"{side} 表头必须唯一匹配；候选：{positions}。请设置 sheets.{side} 和 header_rows.{side}。")
    name, row, mapping, ambiguities = candidates[0]
    if ambiguities:
        raise ReconciliationError(f"{side} {name}!{row} 字段对应多列：{','.join(ambiguities)}；请收窄 aliases。")
    if len(set(mapping.values())) != len(mapping):
        raise ReconciliationError(f"{side} 表头同一列被映射多个字段，请修正 aliases。")
    return name, row, mapping


def _read_fee(formula, cached, raw_values, row_number, mapping):
    if not isinstance(formula, str) or not formula.startswith("="):
        return _money(formula), "原收费数值", "" if _money(formula) is not None else "快递收费缺失或非数值"
    expression = re.sub(r"\s+", "", formula[1:])
    refs = re.fullmatch(r"\$?([A-Za-z]+)\$?(\d+)\+\$?([A-Za-z]+)\$?(\d+)", expression)
    if refs:
        c1, r1, c2, r2 = refs.groups()
        columns = [column_index_from_string(c1), column_index_from_string(c2)]
        expected = [mapping.get("fee_part_1"), mapping.get("fee_part_2")]
        if int(r1) == row_number == int(r2) and None not in expected and sorted(columns) == sorted(expected):
            parts = [_decimal(raw_values[col - 1]) for col in columns]
            if all(value is not None for value in parts):
                amount = _money(sum(parts, ZERO))
                cache = _money(cached)
                note = "本行已确认费用列加法重算；缓存缺失" if cache is None else "本行已确认费用列加法重算；缓存不一致" if cache != amount else "本行已确认费用列加法重算；缓存已核对"
                return amount, note, ""
    return None, "公式未重算", "快递收费公式不属于已确认本行费用列加法，无法确定收费"


def _load_bill(path, side, profile):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True) if side == "courier" else None
    records, excluded, unidentified = [], [], []
    try:
        name, header_row, mapping = _select_sheet(book, side, profile)
        sheet = book[name]
        frows = sheet.iter_rows(min_row=header_row + 1, values_only=True)
        crows = cached[name].iter_rows(min_row=header_row + 1, values_only=True) if cached else None
        for row_number, values in enumerate(frows, header_row + 1):
            cached_values = next(crows) if crows else None
            # System freight has no mapped field and is never extracted or interpreted.
            selected = {key: values[column - 1] if column <= len(values) else None for key, column in mapping.items()}
            source = {"来源": side, "工作表": name, "源行": row_number}
            meaningful_values = list(selected.values()) if side == "system" else values
            if not any(_text(value) for value in meaningful_values):
                excluded.append({**source, "原因": "整行空白"})
                continue
            id_value = selected.get("id")
            first = next((_text(v) for v in meaningful_values if _text(v)), "")
            if _text(id_value) in profile["summary_labels"] or (not _text(id_value) and first in profile["summary_labels"]):
                excluded.append({**source, "原因": "明确汇总标签", "标签": _text(id_value) or first})
                continue
            tracking, id_issue = _identifier(id_value)
            weight = _decimal(selected.get("weight"))
            weight = weight if weight is not None and weight >= 0 else None
            date = _date(selected.get("date"), book.epoch)
            item = {"id": tracking, "row": row_number, "sheet": name, "date": date, "date_raw": _text(selected.get("date")), "weight": weight, "weight_raw": _text(selected.get("weight")), "fields": mapping}
            if side == "courier":
                fee_column = mapping.get("fee")
                formula = selected.get("fee")
                cache = cached_values[fee_column - 1] if fee_column and fee_column <= len(cached_values) else None
                fee, fee_source, fee_issue = _read_fee(formula, cache, values, row_number, mapping)
                item.update(address=_text(selected.get("address")), fee=fee, fee_source=f"{name}!{get_column_letter(fee_column)}{row_number}：{fee_source}" if fee_column else fee_source, fee_issue=fee_issue)
            else:
                product_value = selected.get("product")
                quantity_raw = selected.get("quantity")
                quantity = _decimal(quantity_raw)
                quantity_note = ""
                if _text(quantity_raw) and quantity is None:
                    quantity = quantity_raw
                    quantity_note = "商品件数为非数值或组合文本，原样保留未拆分，不参与核价"
                elif quantity is not None and quantity < 0:
                    quantity_note = "商品件数为负数，需核实；原值保留，不参与核价"
                item.update(address=" / ".join(_text(selected.get(k)) for k in ("province", "city", "district") if _text(selected.get(k))), product=str(product_value) if _text(product_value) else "", quantity=quantity, quantity_note=quantity_note, carrier=_text(selected.get("carrier")))
            if id_issue:
                unidentified.append({**source, "来源键": f"{side}:{name}:{row_number}", "单号原值": _text(id_value), "原因": id_issue, "商品编码": item.get("product", "")})
            else:
                records.append(item)
        metadata = {"file": path.name, "sheet": name, "header_row": header_row, "fields": mapping, "data_rows": len(records) + len(unidentified), "excluded_rows": len(excluded)}
        return records, excluded, unidentified, metadata
    finally:
        book.close()
        if cached:
            cached.close()


def _area(value):
    text = _text(value)
    province = next((p for p in PROVINCES if text.startswith(p)), None)
    if province == "广东":
        remainder = re.sub(r"^广东(?:省)?[\s/]*", "", text)
        city = re.match(r"([^\s/]+?市)", remainder)
        if remainder.startswith("深圳"):
            return "深圳", True
        return "广东", bool(city)
    return province, True if province else False


def _price_area(value):
    if _text(value) in ("深圳", "深圳市"):
        return "深圳"
    return _area(value)[0] or _text(value)


TITLE_RE = re.compile(r"(?:(\d{4})年)?(\d{1,2})月(\d{1,2})日\s*[-—–~～至]\s*(?:(\d{1,2})月)?(\d{1,2})日(?:价格|报价|价格表)?")


def _interval(value):
    text = re.sub(r"\s+", "", _text(value)).lower().replace("公斤", "kg").replace("千克", "kg")
    match = re.fullmatch(r"(\d+(?:\.\d+)?)(?:kg)?[-—–~～](\d+(?:\.\d+)?)(?:kg)", text)
    if match:
        low, high = map(D, match.groups())
        if low < high:
            return low, high
    return None


def _load_prices(path, dates, profile):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cache = openpyxl.load_workbook(path, read_only=True, data_only=True)
    issues, public_rules, versions = [], [], []
    try:
        explicit = profile.get("sheets", {}).get("price")
        if explicit and explicit not in book.sheetnames:
            raise ReconciliationError(f"price 指定工作表不存在：{explicit}")
        candidates = []
        for name in [explicit] if explicit else book.sheetnames:
            rows = list(book[name].iter_rows(values_only=True))
            starts = [(i + 1, col + 1, TITLE_RE.fullmatch(_text(value))) for i, row in enumerate(rows) for col, value in enumerate(row) if TITLE_RE.fullmatch(_text(value))]
            if starts:
                candidates.append((name, rows, starts))
        if len(candidates) != 1:
            raise ReconciliationError("价格表必须唯一匹配 dated_matrix_v1 日期分段矩阵；请设置 sheets.price 或提供支持的价格表结构。")
        name, rows, starts = candidates[0]
        cache_rows = list(cache[name].iter_rows(values_only=True))
        years = {date.year for date in dates if date}
        months = {date.month for date in dates if date}
        inferred_year = next(iter(years)) if len(years) == 1 else None
        for index, (title_row, title_col, match) in enumerate(starts):
            title = _text(rows[title_row - 1][title_col - 1])
            year, month, lowday, endmonth, highday = match.groups()
            valid_year = int(year) if year else inferred_year if months == {int(month)} else None
            start = end = None
            if valid_year:
                try:
                    start = dt.date(valid_year, int(month), int(lowday))
                    end = dt.date(valid_year, int(endmonth or month), int(highday))
                    if end < start:
                        start = end = None
                except ValueError:
                    pass
            version = {"id": f"V{index + 1}", "title": title, "start": start, "end": end, "rows": {}, "columns": [], "blockers": [], "source": f"{name}!{get_column_letter(title_col)}{title_row}"}
            if not start:
                issues.append({"问题类型": "日期范围无法确定", "来源": version["source"], "说明": "日期无有效明确年份，或账单不满足唯一年度且同月，或日期范围无效", "版本": title})
            elif not year:
                issues.append({"问题类型": "年份由账单补足", "来源": version["source"], "说明": f"按两账单唯一年度{valid_year}年及同月记录补足", "版本": title, "严重程度": "说明"})
            stop = starts[index + 1][0] - 1 if index + 1 < len(starts) else len(rows)
            header_candidates = []
            for r in range(title_row + 1, min(stop, title_row + 10) + 1):
                for c, value in enumerate(rows[r - 1], 1):
                    if _header(value) in {_header(x) for x in profile["aliases"]["price_region"]}:
                        header_candidates.append((r, c))
            if len(header_candidates) != 1:
                version["blockers"].append("地区/重量表头无法唯一识别")
                versions.append(version)
                issues.append({"问题类型": "价格表头待确认", "来源": version["source"], "说明": version["blockers"][0], "版本": title})
                continue
            header_row, region_col = header_candidates[0]
            for c, value in enumerate(rows[header_row - 1], 1):
                if c == region_col or not _text(value):
                    continue
                interval = _interval(value)
                column = {"col": c, "label": _text(value), "interval": interval, "source": f"{name}!{get_column_letter(c)}{header_row}"}
                version["columns"].append(column)
                if interval is None:
                    issues.append({"问题类型": "重量范围未定义", "来源": column["source"], "说明": "未从重量原文获得完整上下界；不擅自扩展末档", "版本": title})
            known_intervals = sorted((c["interval"] for c in version["columns"] if c["interval"]), key=lambda v: (v[0], v[1]))
            if any(low >= 3 for low, high in known_intervals):
                issues.append({"问题类型": ">3kg边界开闭未定义", "来源": f"{name}!{header_row}", "说明": "仅严格档内核价；精确相邻端点待确认，同价可保留辅助金额", "版本": title})
            for left, right in zip(known_intervals, known_intervals[1:]):
                if left[1] != right[0]:
                    issues.append({"问题类型": "重量范围冲突" if left[1] > right[0] else "重量范围缺口", "来源": f"{name}!{header_row}", "说明": f"{left} 与 {right}", "版本": title})
            if not version["columns"]:
                version["blockers"].append("无重量列")
            for r in range(header_row + 1, stop + 1):
                row = rows[r - 1]
                raw_region = row[region_col - 1] if region_col <= len(row) else None
                if not _text(raw_region):
                    note = "；".join(_text(v) for v in row if _text(v))
                    if note:
                        version["blockers"].append("存在尚未支持的附注/附加费条款")
                        issues.append({"问题类型": "附注待确认", "来源": f"{name}!{r}", "说明": note, "版本": title})
                    continue
                if _text(raw_region) in profile["summary_labels"]:
                    continue
                region = _price_area(raw_region)
                if region not in PROVINCES and region != "深圳":
                    version["blockers"].append("存在未识别地区或附注")
                    issues.append({"问题类型": "地区/附注待确认", "来源": f"{name}!{get_column_letter(region_col)}{r}", "说明": _text(raw_region), "版本": title})
                    continue
                quotes = []
                text_rule = any(isinstance(row[c["col"] - 1], str) and not row[c["col"] - 1].startswith("=") and re.search(r"首重|续重|加收|附加|续费", row[c["col"] - 1]) for c in version["columns"] if c["col"] <= len(row))
                previous = None
                for column in version["columns"]:
                    c = column["col"]
                    raw = row[c - 1] if c <= len(row) else None
                    value, issue, note = _decimal(raw), "", ""
                    source = f"{name}!{get_column_letter(c)}{r}"
                    if isinstance(raw, str) and raw.startswith("="):
                        value = _decimal(cache_rows[r - 1][c - 1])
                        note = "价格公式缓存未经可靠重算验证，不能作为确定报价"
                        issue = "价格公式待确认（缓存未验证）" if value is not None else "价格公式缓存缺失或非数值"
                        issues.append({"问题类型": "价格公式缓存", "来源": source, "说明": issue or note, "版本": title})
                    if text_rule:
                        issue = "首续重/文字报价参数或适用范围待确认"
                    elif value is None:
                        issue = issue or "价格空白或非数值"
                    elif value < 0:
                        issue = "负价格待确认"
                    elif value == 0:
                        issues.append({"问题类型": "零价格疑点", "来源": source, "说明": "保留原价，不凭常识改价", "版本": title, "地区": region, "严重程度": "疑点"})
                    if issue and (not text_rule or _text(raw)):
                        issues.append({"问题类型": issue, "来源": source, "说明": _text(raw), "版本": title, "地区": region})
                    if value is not None and previous is not None and value < previous:
                        issues.append({"问题类型": "重量增大价格下降", "来源": source, "说明": "保留原价，需核报价约定", "版本": title, "地区": region, "严重程度": "疑点"})
                    if value is not None:
                        previous = value
                    rule_id = f"{version['id']}-{r}-{c}"
                    quotes.append({**column, "value": value, "issue": issue, "note": note, "source": source, "id": rule_id})
                    if not text_rule or _text(raw):
                        public_rules.append({"规则ID": rule_id, "版本": title, "生效起": start.isoformat() if start else None, "生效止": end.isoformat() if end else None, "地区": region, "重量原文": "文字报价范围待确认" if text_rule else column["label"], "基础价": value, "原文": _text(raw), "来源": source, "计价类型": "文字规则" if text_rule else "整票阶梯价", "规则状态": "待确认" if issue or column["interval"] is None else "明确（边界另审）", "说明": issue or note})
                version["rows"].setdefault(region, []).append(quotes)
                if len(version["rows"][region]) > 1:
                    issues.append({"问题类型": "同地区重复报价", "来源": f"{name}!{r}", "说明": region, "版本": title})
            versions.append(version)
        valid = sorted((v for v in versions if v["start"]), key=lambda v: (v["start"], v["end"]))
        for left, right in zip(valid, valid[1:]):
            if right["start"] <= left["end"]:
                issues.append({"问题类型": "价格版本日期重叠", "来源": left["source"] + "；" + right["source"], "说明": "重叠日期禁止任选版本"})
            elif right["start"] > left["end"] + dt.timedelta(days=1):
                issues.append({"问题类型": "价格版本日期缺口", "来源": left["source"] + "；" + right["source"], "说明": f"{left['end']} 至 {right['start']} 中间缺日期"})
        return versions, public_rules, issues, {"file": path.name, "sheet": name, "versions": len(versions)}
    finally:
        book.close()
        cache.close()


def _version(date, versions):
    matches = [v for v in versions if date and v["start"] and v["start"] <= date <= v["end"]]
    return matches[0] if len(matches) == 1 else None


def _quote(weight, version, area):
    result = {"amount": None, "label": "", "source": "", "issue": "", "note": "", "clear": False}
    if weight is None:
        result["issue"] = "重量缺失、非数值或负数"
        return result
    if version is None:
        result["issue"] = "日期缺少唯一价格版本"
        return result
    if version["blockers"]:
        result["issue"] = "；".join(sorted(set(version["blockers"])))
        return result
    rows = version["rows"].get(area, [])
    if len(rows) != 1:
        result["issue"] = "地区缺少报价" if not rows else "同地区重复报价规则冲突"
        return result
    quotes = rows[0]
    candidates = []
    boundary = False
    unsupported = False
    for quote in quotes:
        interval = quote["interval"]
        if interval is None:
            continue
        low, high = interval
        if weight <= 3:
            known = LOW_INTERVALS.get(interval)
            if known and known[0] <= weight <= known[1]:
                candidates.append(quote)
            elif not known and low < 3 and low <= weight <= high:
                candidates.append(quote)
                unsupported = True
        elif low < weight < high:
            candidates.append(quote)
        elif low == weight or high == weight:
            candidates.append(quote)
            boundary = True
    result["label"] = " / ".join(q["label"] for q in candidates)
    result["source"] = "；".join(q["source"] for q in candidates)
    if weight <= 3 and _band(weight) == "重量边界待确认":
        result["issue"] = "重量边界待确认"
    elif unsupported:
        result["issue"] = "价格表低重量档与已确认六档不兼容"
    elif not candidates:
        result["issue"] = "末档重量范围不明或超出价格表覆盖" if any(q["interval"] is None for q in quotes) else "重量未覆盖价格表"
    elif any(q["issue"] for q in candidates):
        result["issue"] = "；".join(sorted({q["issue"] for q in candidates if q["issue"]}))
    elif boundary:
        contiguous = len(candidates) == 2 and candidates[0]["interval"][1] == candidates[1]["interval"][0] == weight
        amounts = {_money(q["value"]) for q in candidates}
        if contiguous and len(amounts) == 1:
            result["amount"] = amounts.pop()
            result["note"] = "边界开闭待确认；相邻档同价，金额唯一"
        else:
            result["issue"] = "价格表>3kg边界开闭不明"
    elif len(candidates) > 1:
        result["issue"] = "重量报价规则冲突"
    else:
        result["amount"] = _money(candidates[0]["value"])
        result["note"] = candidates[0]["note"]
        result["clear"] = True
    return result


CAUSE_ACTIONS = {
    "关联待核实": ("仓库＋快递", "核对重复单号和商品多行关系，确定唯一包裹及收费记录"),
    "仅快递有单": ("仓库", "核实发货和运单归属，并补系统记录"),
    "仅系统有单": ("快递", "带发货凭证核查快递账期及账单覆盖"),
    "地区和日期待确认": ("快递", "确认计费地区与选价日期后补算"),
    "地区待确认": ("快递", "确认计费地区，并核对系统完整收货地址"),
    "日期待确认": ("快递", "确认选价日期及对应价格版本"),
    "资料待补": ("仓库＋快递", "按未核价原因补充重量、日期或账单收费资料"),
    "报价规则待确认": ("快递", "确认缺失或冲突的报价规则后补算"),
    "重量影响运费": ("仓库", "复称含包装重量并核对件数，确认后向快递举证"),
    "计价不符": ("快递", "按快递重量核对报价与账单收费项目"),
    "重量＋计价差异": ("仓库＋快递", "仓库复称重量和件数，快递核对报价及收费项目"),
    "跨档但运费不变": ("仓库", "次要核对重量档，当前无金额争议"),
}


def _joined_source(rows):
    return "；".join(f"{r['sheet']}!{r['row']}" for r in rows)


def _make_record(tracking, courier_rows, system_rows, versions):
    c = courier_rows[0] if len(courier_rows) == 1 else None
    s = system_rows[0] if len(system_rows) == 1 else None
    duplicate = len(courier_rows) > 1 or len(system_rows) > 1
    if len(system_rows) > 1:
        group_type, product = "商品关联待核实", "商品关联待核实"
    elif s and s["product"]:
        group_type, product = "商品", s["product"]
    elif s:
        group_type, product = "商品编码缺失", "商品编码缺失"
    else:
        group_type, product = "仅快递有单·商品待关联", "仅快递有单·商品待关联"
    cw, sw = c["weight"] if c else None, s["weight"] if s else None
    cb, sb = _band(cw), _band(sw)
    ca, city_known = _area(c["address"]) if c else (None, False)
    sa, system_city_known = _area(s["address"]) if s else (None, False)
    area, area_source = ca, "快递目的地"
    if not ca and sa:
        area, area_source, city_known = sa, "系统补足", system_city_known
    cv = _version(c["date"], versions) if c else None
    sv = _version(s["date"], versions) if s else None
    common, system_only = [], []
    region_pending = bool(c and ((ca and sa and ca != sa) or not area or (area == "广东" and not city_known)))
    if c and ca and sa and ca != sa:
        common.append("双方计价地区冲突待核实")
    elif c and not area:
        common.append("目的地区域无法识别")
    elif c and area == "广东" and not city_known:
        common.append("广东城市信息不足，无法排除深圳专价")
    date_pending = False
    if c and not cv:
        common.append("快递日期缺少唯一价格版本")
        date_pending = True
    if c and s and s["date"] is not None and (not sv or not cv or sv["id"] != cv["id"]):
        common.append("双方日期跨价格版本，日期口径待确认")
        date_pending = True
    elif c and s and s["date"] is None:
        system_only.append("系统日期缺失或无效，无法核对是否跨版本")
        date_pending = True
    row = {
        "物流单号": tracking, "商品编码": product, "分组类型": group_type,
        "商品件数": s["quantity"] if s else None, "商品资料缺失": "；".join((["商品编码缺失"] if s and not s["product"] else []) + (["商品件数缺失"] if s and s["quantity"] is None else [])),
        "商品件数核查说明": s["quantity_note"] if s else "",
        "仅快递": bool(courier_rows and not system_rows), "仅系统": bool(system_rows and not courier_rows),
        "匹配结果": "关联待核实" if duplicate else "已匹配" if c and s else "未匹配",
        "快递源工作表": c["sheet"] if c else courier_rows[0]["sheet"] if courier_rows else "",
        "系统源工作表": s["sheet"] if s else system_rows[0]["sheet"] if system_rows else "",
        "快递源行": c["row"] if c else None, "系统源行": s["row"] if s else None,
        "快递源行列表": [r["row"] for r in courier_rows], "系统源行列表": [r["row"] for r in system_rows],
        "快递来源": _joined_source(courier_rows), "系统来源": _joined_source(system_rows),
        "快递日期": c["date"].isoformat() if c and c["date"] else None,
        "系统日期": s["date"].isoformat() if s and s["date"] else None,
        "快递目的地": c["address"] if c else "", "系统省市区": s["address"] if s else "",
        "快递计价地区": ca or "", "系统计价地区": sa or "", "计价地区": area or "", "地区来源": area_source,
        "价格版本": cv["title"] if cv else "", "计价日期": c["date"].isoformat() if c and c["date"] else None, "日期来源": "快递账单日期",
        "快递重量kg": cw, "系统重量kg": sw, "快递重量原值": c["weight_raw"] if c else "", "系统重量原值": s["weight_raw"] if s else "",
        "重量差kg": cw - sw if cw is not None and sw is not None else None, "快递重量区间": cb if c else "未匹配", "系统重量区间": sb if s else "未匹配",
        "原始重量核对": "一致" if cw is not None and sw is not None and cw == sw else "不一致" if cw is not None and sw is not None else "不可比",
        "区间核对": "不可比" if not c or not s else "数据异常或待确认" if "待确认" in cb + sb or "异常" in cb + sb else "一致" if cb == sb else "不一致",
        "快递账单收费": c["fee"] if c else None, "按快递重量计算": None, "按系统重量计算": None,
        "运费来源": c["fee_source"] if c else "", "差额": None, "总差额": None, "重量影响金额": None, "计价差异金额": None,
        "快递计价档": "", "系统计价档": "", "快递价格来源": "", "系统价格来源": "",
        "快递价格状态": "未核价", "系统价格状态": "未核价", "未核价原因": "", "规则备注": "",
        "主原因": "", "首找对象": "", "处理要求": "", "判断依据": "", "重量跨档": False,
        "异常": False, "举证可用": False, "同档小差已忽略": False,
        "系统多行商品": [{"源行": r["row"], "商品编码": r["product"], "商品件数": r["quantity"], "商品件数核查说明": r["quantity_note"]} for r in system_rows] if len(system_rows) > 1 else [],
    }
    failures, notes, clear = [], [], {}
    for prefix, record, weight in (("快递", c, cw), ("系统", s, sw)):
        reasons = list(common)
        if prefix == "系统":
            reasons += system_only
        if not record:
            reasons.append("关联多行待核实" if (courier_rows if prefix == "快递" else system_rows) else "单号未匹配")
        if prefix == "系统" and (not c or duplicate):
            reasons.append("缺少唯一快递关联记录")
        if reasons:
            failures.append(prefix + "：" + "；".join(reasons))
            continue
        quoted = _quote(weight, cv, area)
        row[prefix + "计价档"] = quoted["label"]
        row[prefix + "价格来源"] = quoted["source"]
        clear[prefix] = quoted["clear"]
        if quoted["issue"]:
            failures.append(prefix + "：" + quoted["issue"])
        if quoted["note"]:
            notes.append(prefix + "：" + quoted["note"])
        row["按" + prefix + "重量计算"] = quoted["amount"]
        if quoted["amount"] is not None and row["快递账单收费"] is not None:
            row[prefix + "价格状态"] = "一致" if quoted["amount"] == row["快递账单收费"] else "差异"
    if c and c["fee_issue"]:
        failures.append(c["fee_issue"])
    row["未核价原因"], row["规则备注"] = "；".join(failures), "；".join(notes)
    actual, courier_price, system_price = (row[k] for k in ("快递账单收费", "按快递重量计算", "按系统重量计算"))
    if actual is not None and system_price is not None:
        row["差额"] = row["总差额"] = actual - system_price
    if courier_price is not None and system_price is not None:
        row["重量影响金额"] = courier_price - system_price
    if actual is not None and courier_price is not None:
        row["计价差异金额"] = actual - courier_price
    weight_effect, pricing_effect = row["重量影响金额"], row["计价差异金额"]
    actual_band_change = bool(row["快递计价档"] and row["系统计价档"] and row["快递计价档"] != row["系统计价档"])
    coarse_change = row["区间核对"] == "不一致"
    row["重量跨档"] = coarse_change or actual_band_change
    row["同档小差已忽略"] = bool(c and s and cw is not None and sw is not None and cw != sw and cb == sb and not actual_band_change and courier_price is not None and system_price is not None and courier_price == system_price)
    if duplicate:
        cause = "关联待核实"
    elif row["仅快递"]:
        cause = "仅快递有单"
    elif row["仅系统"]:
        cause = "仅系统有单"
    elif region_pending and date_pending:
        cause = "地区和日期待确认"
    elif region_pending:
        cause = "地区待确认"
    elif date_pending:
        cause = "日期待确认"
    elif weight_effect is None or pricing_effect is None:
        cause = "资料待补" if cw is None or sw is None or actual is None else "报价规则待确认"
    elif notes and any("待确认" in note for note in notes):
        cause = "报价规则待确认"
    elif weight_effect and pricing_effect:
        cause = "重量＋计价差异"
    elif weight_effect:
        cause = "重量影响运费"
    elif pricing_effect:
        cause = "计价不符"
    elif row["重量跨档"]:
        cause = "跨档但运费不变"
    else:
        cause = ""
    row["主原因"] = row["异常说明"] = cause
    row["异常"] = bool(cause)
    if cause:
        row["首找对象"], row["处理要求"] = CAUSE_ACTIONS[cause]
    if duplicate:
        if len(courier_rows) > 1 and len(system_rows) <= 1:
            row["首找对象"] = "快递"
        elif len(system_rows) > 1 and len(courier_rows) <= 1:
            row["首找对象"] = "仓库"
    if row["差额"] is not None and weight_effect is not None and pricing_effect is not None:
        row["判断依据"] = f"重量影响{weight_effect:+.2f}元＋计价差异{pricing_effect:+.2f}元＝总差额{row['差额']:+.2f}元"
    else:
        row["判断依据"] = row["未核价原因"] or cause
    row["排查建议"] = row["处理要求"]
    row["举证可用"] = bool(c and s and not duplicate and row["差额"] is not None and row["差额"] > 0 and courier_price is not None and clear.get("快递") and clear.get("系统") and not failures)
    return row


def _optional_sum(values):
    known = [value for value in values if value is not None]
    return sum(known, ZERO) if known else None


def _product_summaries(details):
    grouped = collections.defaultdict(list)
    for row in details:
        grouped[(row["分组类型"], row["商品编码"])].append(row)
    products = []
    for (group_type, product), rows in grouped.items():
        priced = [r for r in rows if r["差额"] is not None]
        causes = collections.Counter(r["主原因"] for r in rows)
        causes = dict(sorted(causes.items(), key=lambda pair: (-pair[1], list(CAUSE_ACTIONS).index(pair[0]))))
        targets = {target for r in rows for target in r["首找对象"].split("＋") if target}
        target = "＋".join(x for x in ("仓库", "快递") if x in targets)
        total = _optional_sum(r["差额"] for r in rows)
        high = sum((r["差额"] for r in priced if r["差额"] > 0), ZERO) if priced else None
        low = sum((-r["差额"] for r in priced if r["差额"] < 0), ZERO) if priced else None
        actions = list(dict.fromkeys(r["处理要求"] for r in sorted(rows, key=lambda r: list(CAUSE_ACTIONS).index(r["主原因"]))))
        products.append({"商品编码": product, "分组类型": group_type, "异常总单量": len(rows), "总差额": total, "净差额": total, "账单偏高金额": high, "账单偏低金额": low, "未知差额单量": len(rows) - len(priced), "原因单量": causes, "主要异常": "；".join(f"{cause}{count}单" for cause, count in causes.items()), "异常原因（单量）": "；".join(f"{cause}{count}单" for cause, count in causes.items()), "首找对象": target, "找谁处理": target, "首找对象单量": dict(sorted(collections.Counter(r["首找对象"] for r in rows).items())), "处理要求": "；".join(actions), "核查重点": "；".join(actions), "重量影响金额": _optional_sum(r["重量影响金额"] for r in priced), "计价差异金额": _optional_sum(r["计价差异金额"] for r in priced), "已知辅助计价差异金额": _optional_sum(r["计价差异金额"] for r in rows if r["差额"] is None)})
        products[-1]["可分解差额单量"] = sum(r["差额"] is not None and r["重量影响金额"] is not None and r["计价差异金额"] is not None for r in rows)
        products[-1]["可分解差额"] = _optional_sum(r["差额"] for r in priced if r["重量影响金额"] is not None and r["计价差异金额"] is not None)
    return sorted(products, key=lambda p: (-(p["账单偏高金额"] or ZERO), -p["异常总单量"], p["商品编码"], p["分组类型"]))


def _validate_profile(profile):
    allowed = set(default_profile())
    unknown = set(profile) - allowed
    if unknown:
        raise ReconciliationError("不支持的配置项：" + ", ".join(sorted(unknown)))
    if set(profile) != allowed:
        raise ReconciliationError("配置缺少必需项：" + ", ".join(sorted(allowed - set(profile))))
    if profile["profile_version"] != ENGINE_VERSION or profile["profile_id"] != "yuantong-jushuitan-v1":
        raise ReconciliationError("配置规则版本或profile_id与当前引擎不兼容。")
    if not isinstance(profile["config_revision"], str) or not profile["config_revision"].strip():
        raise ReconciliationError("config_revision 必须为非空文本，用于标明字段映射修订。")
    for group, fields in (("sheets", {"courier", "system", "price"}), ("header_rows", {"courier", "system"})):
        if not isinstance(profile[group], dict) or set(profile[group]) - fields:
            raise ReconciliationError(f"不支持的配置分组 {group}，有效键为 {sorted(fields)}。")
    for value in profile["header_rows"].values():
        if value is not None and (type(value) is not int or value < 1):
            raise ReconciliationError("header_rows 必须为空或正整数。")
    if any(value is not None and (not isinstance(value, str) or not value) for value in profile["sheets"].values()):
        raise ReconciliationError("sheets 必须为空或非空工作表名称字符串。")
    if type(profile["header_search_rows"]) is not int or profile["header_search_rows"] < 1:
        raise ReconciliationError("header_search_rows 必须为正整数。")
    base_aliases = default_profile()["aliases"]
    if set(profile["aliases"]) != set(base_aliases):
        raise ReconciliationError("aliases 配置分组不完整或含不支持的项。")
    for side in ("courier", "system"):
        if set(profile["aliases"][side]) != set(base_aliases[side]):
            raise ReconciliationError(f"aliases.{side}字段集合不得更改；可在对应数组调整别名。")
        for aliases in profile["aliases"][side].values():
            if not isinstance(aliases, list) or not aliases or any(not isinstance(alias, str) or not alias.strip() for alias in aliases):
                raise ReconciliationError("字段别名必须为非空字符串数组。")
    for key, values in (("summary_labels", profile["summary_labels"]), ("aliases.price_region", profile["aliases"]["price_region"])):
        if not isinstance(values, list) or not values or any(not isinstance(value, str) or not value.strip() for value in values):
            raise ReconciliationError(f"{key} 必须为非空字符串数组。")
    required = {"price_mode": "dated_matrix_v1", "currency": "CNY", "rounding": "ROUND_HALF_UP", "weight_unit": "kg", "date_policy": "courier_date_cross_version_pending", "region_policy": "province_prefix_city_exception_pending", "duplicate_policy": "pending_no_aggregation", "carrier_policy": "yuantong_brand_aliases", "formula_mode": "same_row_fee_parts_addition", "boundary_policy": "user_six_bands_then_strict_interior"}
    for key, value in required.items():
        if profile.get(key) != value:
            raise ReconciliationError(f"不支持配置 {key}={profile.get(key)!r}；本引擎要求 {value!r}，不能静默忽略规则修改。")
    if set(profile["aliases"]["system"]) - {"id", "weight", "date", "product", "quantity", "carrier", "province", "city", "district"}:
        raise ReconciliationError("系统字段映射包含不支持的字段；禁止映射系统原运费。")
    if any("运费" in alias or "freight" in alias.lower() for aliases in profile["aliases"]["system"].values() for alias in aliases):
        raise ReconciliationError("禁止将系统原运费映射到任何核对字段。")
    if profile.get("city_exceptions") != {"广东": ["深圳"]}:
        raise ReconciliationError("本规则版本仅支持已确认的广东/深圳城市专价结构。")


def _reconcile(courier, system, price, profile):
    _validate_profile(profile)
    crows, ce, cu, cm = _load_bill(courier, "courier", profile)
    srows, se, su, sm = _load_bill(system, "system", profile)
    carriers = sorted({row["carrier"] for row in srows if row["carrier"]})
    carrier_pattern = re.compile(r"(?:[\u4e00-\u9fffA-Za-z0-9（）() -]*[-—]?)?圆通(?:速递|快递)(?:_\d+)?")
    unsupported_carriers = [carrier for carrier in carriers if not carrier_pattern.fullmatch(carrier)]
    if unsupported_carriers:
        raise ReconciliationError("当前配置只支持圆通品牌及明确网点别名，请按承运商分批并使用对应配置；无法识别：" + "、".join(unsupported_carriers))
    sm["carrier_values"] = carriers
    sm["carrier_policy"] = profile["carrier_policy"]
    sm["carrier_brand"] = "圆通" if carriers else "未提供"
    sm["scope_note"] = "范围以本次提供文件为准；缺少双方主体与快递范围字段时不能证明同主体，单侧缺单需核账期覆盖。"
    versions, price_rules, issues, pm = _load_prices(price, [r["date"] for r in crows + srows], profile)
    cg, sg = collections.defaultdict(list), collections.defaultdict(list)
    for row in crows:
        cg[row["id"]].append(row)
    for row in srows:
        sg[row["id"]].append(row)
    records = [_make_record(tracking, cg.get(tracking, []), sg.get(tracking, []), versions) for tracking in sorted(set(cg) | set(sg))]
    details = [r for r in records if r["异常"]]
    products = _product_summaries(details)
    ordering = {(p["分组类型"], p["商品编码"]): i for i, p in enumerate(products)}
    details.sort(key=lambda r: (ordering[(r["分组类型"], r["商品编码"])], r["物流单号"]))
    evidence = sorted((r for r in records if r["举证可用"]), key=lambda r: (-r["差额"], r["商品编码"], r["物流单号"]))
    priced = [r for r in details if r["差额"] is not None]
    decomposable = [r for r in priced if r["重量影响金额"] is not None and r["计价差异金额"] is not None]
    causes = dict(sorted(collections.Counter(r["主原因"] for r in details).items(), key=lambda p: list(CAUSE_ACTIONS).index(p[0])))
    summary = {
        "承运商品牌": sm["carrier_brand"], "承运商归并策略": profile["carrier_policy"], "承运商原名数量": len(carriers),
        "快递记录行数": len(crows) + len(cu), "系统记录行数": len(srows) + len(su), "快递有效单号记录行数": len(crows), "系统有效单号记录行数": len(srows),
        "快递去重单号数": len(cg), "系统去重单号数": len(sg), "匹配单号数": len(set(cg) & set(sg)), "仅快递有": len(set(cg) - set(sg)), "仅系统有": len(set(sg) - set(cg)),
        "重复单号数": len({key for key in set(cg) | set(sg) if len(cg.get(key, [])) > 1 or len(sg.get(key, [])) > 1}),
        "待补单号记录数": len(cu) + len(su), "排除记录行数": len(ce) + len(se), "全量去重运单数": len(records),
        "异常总单量": len(details), "正常运单数": len(records) - len(details), "涉及商品组数": len(products), "已识别商品数": sum(p["分组类型"] == "商品" for p in products),
        "未知差额单量": len(details) - len(priced), "总差额": sum((r["差额"] for r in priced), ZERO), "净差额": sum((r["差额"] for r in priced), ZERO),
        "账单偏高金额": sum((r["差额"] for r in priced if r["差额"] > 0), ZERO), "账单偏低金额": sum((-r["差额"] for r in priced if r["差额"] < 0), ZERO),
        "重量影响金额": sum((r["重量影响金额"] for r in priced if r["重量影响金额"] is not None), ZERO), "计价差异金额": sum((r["计价差异金额"] for r in priced if r["计价差异金额"] is not None), ZERO),
        "主原因单量": causes, "同档小差忽略单量": sum(r["同档小差已忽略"] for r in records), "商品编码缺失单量": sum("商品编码缺失" in r["商品资料缺失"] for r in details), "商品件数缺失单量": sum("商品件数缺失" in r["商品资料缺失"] for r in details),
        "全量商品编码缺失单量": sum("商品编码缺失" in r["商品资料缺失"] for r in records), "全量商品件数缺失单量": sum("商品件数缺失" in r["商品资料缺失"] for r in records),
        "举证单量": len(evidence), "举证多收金额": sum((r["差额"] for r in evidence), ZERO), "价格规则数": len(price_rules), "价格表问题数": len([i for i in issues if i.get("严重程度") != "说明"]),
    }
    pricing_records = [record for record in records if record["快递源行列表"]]
    summary["核价统计口径"] = "快递账单中可识别的去重单号；两套核价状态使用相同分母，系统独有单号另列"
    summary["核价统计单量"] = len(pricing_records)
    for prefix in ("快递", "系统"):
        for status in ("一致", "差异", "未核价"):
            summary[prefix + "重量核价" + status] = sum(r[prefix + "价格状态"] == status for r in pricing_records)
    affected = collections.Counter(r["未核价原因"] for r in records if r["未核价原因"])
    for reason, count in sorted(affected.items()):
        if "地区缺少报价" in reason or "日期缺少唯一价格版本" in reason or "重量未覆盖" in reason:
            issues.append({"问题类型": "账单覆盖范围缺少规则", "说明": reason, "受影响运单数": count})
    summary["价格表问题数"] = len([issue for issue in issues if issue.get("严重程度") != "说明"])
    summary["可分解差额单量"] = len(decomposable)
    summary["可分解差额"] = _optional_sum(r["差额"] for r in decomposable) if details else ZERO
    summary["重量影响金额"] = _optional_sum(r["重量影响金额"] for r in decomposable) if details else ZERO
    summary["计价差异金额"] = _optional_sum(r["计价差异金额"] for r in decomposable) if details else ZERO
    if details and not priced:
        for key in ("总差额", "净差额", "账单偏高金额", "账单偏低金额", "重量影响金额", "计价差异金额"):
            summary[key] = None
    checks = {
        "单号并集守恒": len(records) == len(cg) + len(sg) - len(set(cg) & set(sg)),
        "快递源行守恒": sum(len(r["快递源行列表"]) for r in records) + len(cu) == cm["data_rows"],
        "系统源行守恒": sum(len(r["系统源行列表"]) for r in records) + len(su) == sm["data_rows"],
        "异常主原因互斥完整": sum(causes.values()) == len(details) == sum(p["异常总单量"] for p in products),
        "异常单号唯一": len({r["物流单号"] for r in details}) == len(details),
        "金额逐单守恒": all(r["差额"] == r["重量影响金额"] + r["计价差异金额"] for r in records if all(r[k] is not None for k in ("差额", "重量影响金额", "计价差异金额"))),
        "可分解集合金额守恒": not decomposable or summary["可分解差额"] == summary["重量影响金额"] + summary["计价差异金额"],
        "正负差额守恒": summary["总差额"] is None or summary["总差额"] == summary["账单偏高金额"] - summary["账单偏低金额"],
        "商品汇总金额守恒": (sum((p["总差额"] for p in products if p["总差额"] is not None), ZERO) if not details or priced else None) == summary["总差额"],
        "举证仅正差且依据明确": all(r["差额"] > 0 and r["举证可用"] for r in evidence),
        "举证单号唯一": len({r["物流单号"] for r in evidence}) == len(evidence),
        "系统原运费未映射": "fee" not in sm["fields"],
    }
    if not all(checks.values()):
        raise AssertionError({name: passed for name, passed in checks.items() if not passed})
    return {"engine_version": ENGINE_VERSION, "profile_version": profile["profile_version"], "profile_id": profile["profile_id"], "config_revision": profile["config_revision"], "summary": summary, "details": details, "all_records": records, "products": products, "evidence": evidence, "price_rules": price_rules, "rule_issues": issues, "excluded_rows": ce + se, "unidentified_records": cu + su, "validation": {"passed": all(checks.values()), "checks": checks}, "input_metadata": {"courier": cm, "system": sm, "price": pm}}
